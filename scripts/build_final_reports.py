#!/usr/bin/env python3
"""Build the four-page submission brief and detailed technical report."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "artifacts" / "figures"
REPORTS = ROOT / "reports"

NAVY = "183153"
BLUE = "2C6ECB"
TEAL = "1C8E8E"
PURPLE = "7655AD"
ORANGE = "D97706"
PALE_BLUE = "EEF4FB"
PALE_GRAY = "F4F6F8"
MID_GRAY = "667485"
LIGHT_BORDER = "D9D9D9"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    node = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), LIGHT_BORDER)
        node.append(e)
    tbl_pr.append(node)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def configure_document(doc: Document, compact: bool = False) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55 if compact else 0.68)
    section.bottom_margin = Inches(0.52 if compact else 0.62)
    section.left_margin = Inches(0.58 if compact else 0.72)
    section.right_margin = Inches(0.58 if compact else 0.72)
    section.header_distance = Inches(0.22)
    section.footer_distance = Inches(0.24)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.25 if compact else 10.2)
    normal.font.color.rgb = RGBColor(30, 41, 53)
    normal.paragraph_format.space_after = Pt(4.0 if compact else 5.0)
    normal.paragraph_format.line_spacing = 1.03 if compact else 1.06

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(24 if compact else 30)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(4)

    subtitle = styles.add_style("Report Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle.font.name = "Aptos"
    subtitle.font.size = Pt(10 if compact else 12)
    subtitle.font.color.rgb = RGBColor.from_string(MID_GRAY)
    subtitle.paragraph_format.space_after = Pt(8)

    for name, size, before, after in (
        ("Heading 1", 16 if compact else 19, 6, 4),
        ("Heading 2", 12 if compact else 14, 6, 3),
        ("Heading 3", 10 if compact else 11.5, 4, 2),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(7.5 if compact else 8.5)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(4)

    source = styles.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
    source.font.name = "Aptos"
    source.font.size = Pt(7.2 if compact else 8)
    source.font.color.rgb = RGBColor.from_string(MID_GRAY)
    source.paragraph_format.space_before = Pt(1)
    source.paragraph_format.space_after = Pt(4)

    small = styles.add_style("Small Body", WD_STYLE_TYPE.PARAGRAPH)
    small.font.name = "Aptos"
    small.font.size = Pt(8.1 if compact else 9)
    small.font.color.rgb = RGBColor(30, 41, 53)
    small.paragraph_format.space_after = Pt(3)
    small.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    header.text = "WG Recommendation Model"
    header.style = styles["Source Note"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.add_run("Afori take home   |   September 2026   |   ")
    add_page_number(footer)


def set_keep(paragraph, keep_next: bool = False) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_lines = OxmlElement("w:keepLines")
    p_pr.append(keep_lines)
    if keep_next:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p = doc.add_paragraph(style="Report Subtitle")
    p.add_run(subtitle)


def add_body(doc: Document, text: str, style: str | None = None, bold_lead: str | None = None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc: Document, items: list[str], compact: bool = False) -> None:
    for item in items:
        p = doc.add_paragraph(style="Small Body" if compact else "Normal")
        p.style = doc.styles["Small Body" if compact else "Normal"]
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.13)
        p.add_run("•  ")
        p.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None, compact: bool = False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 70 if compact else 95, 85, 70 if compact else 95, 85)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(7.6 if compact else 8.8)
        run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Inches(widths[idx])
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for c_idx, text in enumerate(row):
            cell = cells[c_idx]
            if r_idx % 2:
                shade(cell, PALE_GRAY)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 65 if compact else 90, 85, 65 if compact else 90, 85)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(text))
            run.font.name = "Aptos"
            run.font.size = Pt(7.7 if compact else 9)
            if widths:
                cell.width = Inches(widths[c_idx])
    borders(table)
    return table


def add_figure(doc: Document, filename: str, caption: str, width: float, source: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run()
    run.add_picture(str(FIG / filename), width=Inches(width))
    c = doc.add_paragraph(caption, style="Caption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if source:
        s = doc.add_paragraph(source, style="Source Note")
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def set_core_properties(doc: Document, title: str, subject: str) -> None:
    doc.core_properties.title = title
    doc.core_properties.subject = subject
    doc.core_properties.author = "Lakshay Chhabra"
    doc.core_properties.keywords = "Wohngebäude, recommendation, Qwen, LoRA, SFT, GRPO"


def build_brief(path: Path) -> None:
    doc = Document()
    configure_document(doc, compact=True)
    set_core_properties(doc, "WG Recommendation Model Submission Brief", "Four-page take-home report")

    add_title(doc, "WG Recommendation Model Submission Brief", "Qwen3.5 2B LoRA for ranked residential building insurance recommendations")
    p = add_body(
        doc,
        "We trained a 2B open model to rank three offered tariffs as strict JSON. Supervised fine tuning raised held-out Fable Top-1 agreement from 22% to 84%, matching the strongest supplied frontier result, while measured self-hosted serving cost was €0.0028 to €0.0539 per 1,000 recommendations depending on utilization. The selected deliverable is the SFT adapter; a ten-step GRPO continuation did not improve the primary metric.",
        style="Small Body",
    )
    set_keep(p)
    add_table(
        doc,
        ["Selected model", "Golden Top 1", "Top 3 recall", "Strict JSON", "Planning cost per 1k"],
        [["Qwen3.5 2B SFT LoRA", "84%", "86.7%", "100%", "€0.0113"]],
        widths=[1.75, 1.1, 1.15, 1.1, 1.55],
        compact=True,
    )
    doc.add_heading("Experiment design", level=1)
    add_body(
        doc,
        "The held-out benchmark contained 50 Fable-ranked scenarios. It was excluded from training, reward selection, prompt tuning, and checkpoint selection. We generated separate risks, collected real tariff panels through the supplied quote engine, labelled them with GPT-6 Sol, and assigned every scenario exactly one experimental role.",
        style="Small Body",
    )
    add_figure(doc, "01-dataset-funnel-and-splits.png", "Figure 1  Collection funnel and leakage-safe experimental splits", 6.6)
    add_body(
        doc,
        "Of 2,000 priced risks, 1,490 produced at least three quotes. The final data contained 9,632 quotes from nine products and five insurers. Scenario IDs were hashes of normalized risks; exact golden-risk matches were rejected. Quote calls were cached and separated by at least three seconds.",
        style="Small Body",
    )

    add_page_break(doc)
    doc.add_heading("Reward audit and calibration", level=1)
    add_body(
        doc,
        "The supplied value-for-money reward was poorly aligned with the desired quality-first ranking. Deductible and coverage count never varied within a quote panel, and insured amount was effectively neutral in almost every panel under the scorer's noise floor. Coverage score and annual price were the only consistent ranking signals.",
        style="Small Body",
    )
    add_figure(doc, "03-reward-calibration.png", "Figure 2  Reward search results on the 200-scenario calibration split", 6.65)
    add_body(
        doc,
        "A deterministic search evaluated 20,000 configurations on 200 silver-labelled scenarios. Five-fold nested selection chose the same candidate in every fold. It increased teacher Top-1 agreement from 49% to 94%. A permutation control returned chance-level ranking, which argues against label leakage or a mechanically inflated evaluator.",
        style="Small Body",
    )
    add_table(
        doc,
        ["Weight", "Coverage count", "Insured amount", "Deductible", "Coverage score", "Price"],
        [["Selected", "0.0928", "0.4103", "0.0696", "0.4273", "1.7936"]],
        widths=[1.0, 1.1, 1.15, 1.0, 1.1, 1.0],
        compact=True,
    )
    add_body(
        doc,
        "The calibrated scorer later reached 84% Fable Top-1 without model inference. However, its Top-3 recall was 79.3%, compared with 86.7% for SFT: the model recovered 130 of 150 reference positions versus 119 for the scorer. The symbolic result is still a strong baseline, not independent truth: its weights approximate GPT-6 Sol, which itself agreed with Fable on 84% of golden Top-1 choices.",
        style="Small Body",
    )

    add_page_break(doc)
    doc.add_heading("Model training and selection", level=1)
    add_body(
        doc,
        "We selected Qwen3.5-2B because the task is narrow, text-only, and requires a short structured ranking. BF16 LoRA trained 10.9 million parameters, 0.49% of the model. SFT used 300 training and 90 validation examples, completion-only loss, three epochs, and the canonical prompt. The run completed in 6.3 minutes on one RTX A6000 and selected step 21 by validation loss.",
        style="Small Body",
    )
    add_figure(doc, "06-model-progress.png", "Figure 3  Golden model progression with reward calibration shown separately", 6.7)
    add_table(
        doc,
        ["Golden model", "Top 1", "Top 3 recall", "Strict JSON", "Reward ratio", "Mean completion"],
        [
            ["Qwen3.5 2B base", "22%", "66.7%", "2%", "0.844", "493.8 tok"],
            ["SFT selected", "84%", "86.7%", "100%", "0.812", "43.3 tok"],
            ["SFT plus GRPO 10", "84%", "86.0%", "100%", "0.812", "43.3 tok"],
        ],
        widths=[1.35, 0.72, 0.9, 0.9, 0.9, 1.15],
        compact=True,
    )
    doc.add_heading("Why select SFT over the symbolic scorer", level=2)
    add_body(
        doc,
        "The scorer and SFT tied on the first recommendation, but the required deliverable is a ranked Top-3. SFT recovered 11 more of the 150 Fable-selected tariff positions while retaining 100% valid JSON. If production consumed only rank one, the zero-inference-cost scorer would be the simpler choice; for the complete shortlist, SFT has the measured quality advantage.",
        style="Small Body",
    )
    doc.add_heading("Why GRPO stopped at ten steps", level=2)
    add_body(
        doc,
        "GRPO used 810 unlabelled prompts, four sampled completions, and a frozen reward with format, validity, Top-1 utility, and discounted Top-3 utility components. It was stable and moved development Top-1 from 83/90 to 84/90, but reduced development Top-3 recall from 87.0% to 86.3%. On golden it changed no Top-1 decision and reduced Top-3 recall on one scenario. This does not show that longer GRPO cannot help. It shows that ten steps did not establish a broad improvement, while further optimization could amplify errors inherited from the silver teacher and its reward proxy. SFT was therefore the simpler supported choice.",
        style="Small Body",
    )

    add_page_break(doc)
    doc.add_heading("Quality cost and deployment", level=1)
    add_figure(doc, "07-cost-quality.png", "Figure 4  Held-out quality against measured or supplied serving cost", 6.7)
    add_body(
        doc,
        "A vLLM benchmark served the selected adapter on a $0.74 per hour RTX 4090. The sustained 1,200-request run reached 63.66 recommendations per second at concurrency 64, with p50 1.003 seconds, p95 1.105 seconds, 90.7% mean GPU utilization, and no failed or invalid responses. GPU-only cost was €0.002832 per 1,000 at saturation and €0.011327 at the 25% utilization planning assumption. Storage, startup, networking, redundancy, and operational overhead are excluded.",
        style="Small Body",
    )
    doc.add_heading("What the experiment cost", level=2)
    add_body(
        doc,
        "Teacher selection and batch labelling cost approximately $3.80. The total RunPod bill for training, evaluation, and serving benchmarking was $4.00, giving an end-to-end recorded cash spend of approximately $7.80. This is historical experiment spend; it is separate from the per-recommendation production estimate above.",
        style="Small Body",
    )
    doc.add_heading("Conclusions", level=1)
    add_bullets(
        doc,
        [
            "Reward literacy mattered as much as model training. The original reward emphasized price; coverage score recovered the teacher's quality preference.",
            "SFT produced the useful model gain and solved output control. The lower supplied reward ratio is an expected quality-versus-price trade-off, not a contradiction.",
            "GRPO remains plausible but inconclusive. A credible follow-up needs better expert supervision, a new untouched benchmark, and a predeclared development-only checkpoint rule.",
            "The main external-validity risks are the 50-scenario benchmark, synthetic risks, nine-product catalogue, heuristic hazard context, and silver rather than human labels.",
        ],
        compact=True,
    )
    add_body(
        doc,
        "Reproduction package  release model adapters, prepared non-golden data, exact configs, aggregate summaries, hashes, HOW_TO_RUN.md, and scripts/reproduce_model.py. Apple Silicon reproduced all 50 SFT outputs exactly.",
        style="Source Note",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def section_intro(doc: Document, number: str, title: str, summary: str, break_before: bool = True) -> None:
    heading = doc.add_heading(f"{number} {title}", level=1)
    heading.paragraph_format.page_break_before = break_before
    p = add_body(doc, summary)
    p.runs[0].bold = True


def build_technical(path: Path) -> None:
    doc = Document()
    configure_document(doc, compact=False)
    set_core_properties(doc, "WG Recommendation Model Technical Experiment Report", "Detailed experimental record and evidence map")

    add_title(doc, "WG Recommendation Model Technical Experiment Report", "Dataset construction reward calibration Qwen3.5 2B LoRA training GRPO evaluation and serving")
    add_body(doc, "Prepared for the Afori machine learning engineering take home", style="Report Subtitle")
    doc.add_paragraph()
    add_figure(doc, "06-model-progress.png", "Model progression and the separate reward-calibration result", 6.5)
    add_body(
        doc,
        "Main conclusion",
        style="Heading 2",
    )
    add_body(
        doc,
        "A compact SFT LoRA moved Qwen3.5-2B from 22% to 84% Top-1 agreement with the held-out Fable ranking and from 2% to 100% strict-JSON compliance. A ten-step GRPO continuation was stable but did not improve golden Top-1 and slightly reduced Top-3 recall, so the SFT adapter is the selected model. Reward calibration exposed the central task insight: coverage quality and price, rather than the supplied protection-count terms, explained most of the reference ranking.",
    )
    add_table(
        doc,
        ["Artifact", "Selected result"],
        [
            ["Model", "Qwen3.5-2B SFT LoRA"],
            ["Golden Top-1", "42 of 50  84%"],
            ["Golden Top-3 recall", "86.7%"],
            ["Output compliance", "100% strict JSON"],
            ["Measured serving", "63.66 rec/s at concurrency 64"],
            ["Planning cost", "€0.011327 per 1,000 at 25% utilization"],
            ["Recorded experiment spend", "Approximately $7.80 total"],
        ],
        widths=[2.0, 4.55],
    )
    add_body(doc, "Report date  27 September 2026", style="Source Note")

    add_page_break(doc)
    doc.add_heading("Contents", level=1)
    contents = [
        "1 Executive summary",
        "2 Task and evaluation contract",
        "3 Experimental design and data separation",
        "4 Synthetic scenarios and quote collection",
        "5 Silver teacher and dataset splits",
        "6 Reward audit grid search and validation",
        "7 Model selection and RunPod execution",
        "8 Supervised fine tuning",
        "9 GRPO continuation and stopping decision",
        "10 Final evaluation and model selection",
        "11 Serving benchmark and economics",
        "12 Reproducibility and release",
        "13 Limitations and future experiments",
        "Appendix A Decision trail",
        "Appendix B Evidence map and open questions",
    ]
    add_bullets(doc, contents)
    doc.add_heading("Reading guide", level=2)
    add_body(doc, "Sections 1 through 3 define the claim and experimental safeguards. Sections 4 through 6 cover data and reward design. Sections 7 through 10 document model training and evaluation. Sections 11 through 13 address economics, reproducibility, limitations, and future work. The appendices map claims to repository evidence and identify the few points that require owner wording rather than experimental reconstruction.")
    doc.add_heading("Evidence convention", level=2)
    add_body(doc, "Machine-readable summaries, manifests, hashes, and run records control exact numbers. The RunPod journal controls execution chronology. DECISIONS.md controls recorded choices. gpt.md provides narrative context but includes interim plans and is not used alone to support a quantitative claim.")

    section_intro(doc, "1", "Executive summary", "The experiment reached frontier-level agreement with a small open model, but the strongest scientific result was that the task reward itself needed to be understood before reinforcement learning could be useful.")
    add_body(doc, "The supplied benchmark frames the task as ranking three Wohngebäude insurance tariffs for each building and quote panel. The original open-model reference was Qwen3.5-9B at 42% Top-1 agreement. We deliberately chose the smaller Qwen3.5-2B and trained a BF16 LoRA adapter for a narrow structured-output policy.")
    add_body(doc, "The data pipeline produced 1,490 independently generated and quoted scenarios. GPT-6 Sol supplied silver Top-3 labels. Separate pools supported reward calibration, supervised training, and GRPO, while the 50-scenario Fable benchmark remained outside the pipeline until locked evaluation.")
    add_body(doc, "Reward analysis showed that deductible and number of coverages were flat within every quote panel. Under the scorer's noise floor, insured amount also contributed almost no ranking variation. Coverage score and annual price carried the effective signal. A 20,000-candidate search raised agreement with the silver teacher from 49% to 94% Top-1 on the calibration procedure and survived permutation, held-out, and subgroup checks.")
    add_body(doc, "SFT delivered the main learned-model improvement. Golden Top-1 increased from 11/50 to 42/50, Top-3 recall from 66.7% to 86.7%, and strict JSON from 2% to 100%. Ten GRPO steps added one correct Top-1 decision on the silver development set but reduced Top-3 recall there; on golden, Top-1 and reward were unchanged and Top-3 recall declined by one recovered tariff. We therefore selected SFT and retained GRPO as an ablation.")
    add_body(doc, "A measured vLLM run on an RTX 4090 confirmed deployability. At concurrency 64 the adapter served 1,200 requests at 63.66 recommendations per second, p95 latency of 1.105 seconds, and 100% valid output. The GPU-only planning cost is €0.011327 per 1,000 recommendations at 25% utilization. Recorded cash spend for the complete experiment was approximately $7.80: about $3.80 for teacher selection and labels plus a $4.00 RunPod bill for training, evaluation, and serving benchmarks.")

    section_intro(doc, "2", "Task and evaluation contract", "The model receives one building risk, assessed hazard context, and a priced quote panel, then returns exactly three distinct offered tariff IDs in ranked JSON.")
    doc.add_heading("Required output", level=2)
    add_body(doc, '{"top_3":[{"rank":1,"tariff_id":"49"},{"rank":2,"tariff_id":"48"},{"rank":3,"tariff_id":"162"}]}', style="Small Body")
    doc.add_heading("Metrics", level=2)
    add_table(
        doc,
        ["Metric", "Interpretation"],
        [
            ["Top-1 agreement", "Predicted first tariff equals the Fable or split-specific teacher first tariff"],
            ["Top-3 recall", "Fraction of the three reference tariffs recovered in the model ranking"],
            ["Reward ratio", "Supplied reward of the predicted first tariff divided by scenario reward argmax"],
            ["Strict JSON", "Output exactly satisfies the required schema"],
            ["Validity", "All three IDs are distinct and offered in the quote panel"],
        ],
        widths=[1.5, 5.05],
    )
    doc.add_heading("Three kinds of supervision", level=2)
    add_body(doc, "Fable rankings are the 50-scenario evaluation reference supplied with the task. GPT-6 Sol labels are scalable silver supervision and agreed with Fable on 84% of golden Top-1 decisions. The calibrated numeric reward approximates the GPT-6 Sol policy using quote attributes. These three objects are related but not interchangeable.")
    doc.add_heading("Held-out boundary", level=2)
    add_bullets(doc, [
        "Golden labels were never used for training, prompt design, reward selection, or checkpoint selection.",
        "Scenario generation read golden risk fields only to reject exact risk duplicates.",
        "Routine reproduction uses non-golden development data; golden evaluation requires an explicit flag.",
        "After the final SFT and GRPO comparisons were observed, no weights or settings were changed in response.",
    ])
    add_body(doc, "The golden file is small, so differences of one scenario should be reported as observations rather than decisive estimates of population performance.", style="Source Note")

    section_intro(doc, "3", "Experimental design and data separation", "The design assigns each generated scenario one role so that reward search, imitation learning, policy optimization, and evaluation cannot silently share labels.")
    add_figure(doc, "01-dataset-funnel-and-splits.png", "Figure 1  Collection funnel and disjoint split architecture", 6.85)
    add_body(doc, "The first split created a 200-record reward-calibration pool, a 390-record SFT pool, and a 900-record GRPO pool. Independent second-stage splits produced 300/90 SFT and 810/90 GRPO train/validation partitions. Stratification balanced hazard, age, building type, construction class, natural-hazard request, and owner occupancy.")
    add_table(
        doc,
        ["Split", "Records", "Teacher label available", "Purpose"],
        [
            ["Reward calibration", "200", "Yes", "Select and cross-validate reward weights"],
            ["SFT train", "300", "Yes", "Completion-only imitation learning"],
            ["SFT validation", "90", "Yes", "Loss-based checkpoint selection"],
            ["GRPO train", "810", "No", "Policy optimization from frozen numeric reward"],
            ["GRPO validation", "90", "Evaluation only", "Offline comparison and gate"],
            ["Golden", "50", "Fable reference", "Final reporting only"],
        ],
        widths=[1.35, 0.7, 1.45, 3.05],
    )
    doc.add_heading("Canonical prompt", level=2)
    add_body(doc, "Teacher annotation, SFT, GRPO preparation, and evaluation use the same prompt builder. Quote fields are allow-listed, the original numeric reward is excluded from teacher requests, quote order is seeded, and the tariff mapping is stored with each scenario. This prevents prompt drift and reward leakage.")

    section_intro(doc, "4", "Synthetic scenarios and quote collection", "The pipeline generated diverse building risks and priced them through the supplied live engine using a cached rate-limited client.")
    doc.add_heading("Scenario generation", level=2)
    add_body(doc, "Seeded generation created 2,400 unique risks. Scenario IDs are SHA-256 hashes of canonical normalized risk JSON. Structural checks covered enums, construction year, dwelling units, rebuild-sum plausibility, uniqueness, and golden-risk collisions. An early pilot that synthesized postcode suffixes was retired; the final run sampled only known-valid supplied postcodes and classified their hazard tier through the provided risk-context function.")
    doc.add_heading("Quote collection", level=2)
    add_body(doc, "The collector considered 2,000 risks and made 2,000 uncached calls. Each cache record is terminal unless an explicit transient-error retry is requested. Calls were separated by at least three seconds, and only panels with at least three quotes entered the training dataset.")
    add_table(
        doc,
        ["Outcome", "Count", "Share of attempts"],
        [
            ["Successful quote panels", "1,490", "74.5%"],
            ["Insufficient quotes", "332", "16.6%"],
            ["No quotes", "178", "8.9%"],
            ["Errors", "0", "0.0%"],
        ],
        widths=[3.2, 1.3, 2.0],
    )
    add_figure(doc, "02-dataset-distribution.png", "Figure 2  Distribution of the 1,490 usable quoted scenarios", 6.85)
    add_body(doc, "The final dataset contains 9,632 quotes, averaging 6.46 per scenario, from five insurers and nine products. Age, hazard, building type, and supported construction classes remained broad. FHG3 was absent because all 397 attempted FHG3 risks produced fewer than three usable quotes. That absence is a live-engine support limitation and a source of external-validity risk.")

    section_intro(doc, "5", "Silver teacher and dataset splits", "GPT-6 Sol supplied scalable ranking labels, but the experiment treats them as a proxy rather than ground truth.")
    doc.add_heading("Teacher selection", level=2)
    add_body(doc, "The original reward argmax matched Fable on only 38% of golden Top-1 choices. GPT-6 Sol reached 84% Top-1 and 86.7% Top-3 recall on the same benchmark, supporting its use as a silver teacher. That result also defines its limitation: eight of fifty first choices differed from Fable.")
    doc.add_heading("Batch annotation", level=2)
    add_body(doc, "Two offline request shards of 745 scenarios were built from sanitized quote fields and the canonical prompt. Both batches completed with 1,490 valid annotations and zero schema problems. Provider-reported usage was 1,546,926 input tokens and 417,797 output or reasoning tokens. Total annotation cost was $3.635911, or about $0.00244 per labelled scenario.")
    add_table(
        doc,
        ["Control", "Implementation"],
        [
            ["Reward leakage", "Teacher requests exclude reward and recommended-tariff fields"],
            ["Schema", "Structured output validates ranks and offered tariff IDs"],
            ["Metadata", "Insurer and product names are copied from source quotes after validation"],
            ["Traceability", "Request, response, annotation, and source hashes are preserved"],
            ["Resumability", "Prepare submit status download and merge stages persist state"],
        ],
        widths=[1.45, 5.1],
    )
    doc.add_heading("Interpretation", level=2)
    add_body(doc, "SFT learns the silver teacher directly. GRPO does not see teacher labels; it optimizes a numeric reward fitted to a separate subset of those labels. Consequently, policy optimization can improve agreement with the proxy while leaving Fable agreement unchanged or worse. A longer run is scientifically meaningful only with a separate expert-quality development target.")

    section_intro(doc, "6", "Reward audit grid search and validation", "Reward calibration was a constrained search over the supplied scoring formula, not a new model and not a search on golden labels.")
    doc.add_heading("Signal audit", level=2)
    add_body(doc, "Annual price and coverage score varied in every calibration panel. Deductible and number of coverages varied in none. Insured amount had raw variation in 168 of 200 panels but exceeded the scorer's one-percent normalization noise floor in only two. On all 50 golden panels, the frozen ranking is equivalent to a constant protection intercept plus normalized coverage score divided by a price penalty.")
    add_body(doc, "Effective golden formula", style="Heading 3")
    add_body(doc, "reward = (0.2863385 + 0.4273230 × normalized coverage score) ÷ (1 + 1.7936116 × normalized annual price)", style="Small Body")
    doc.add_heading("Search and selection", level=2)
    add_body(doc, "The deterministic search evaluated 20,000 candidates over a simplex for static protection, insured amount, and coverage score, plus a log-uniform price penalty from 0.05 to 5.0 and hand-designed anchors. Candidates were ordered by Top-1, Top-3 recall, pairwise accuracy, and finally distance from the default reward. Reward values were rounded exactly as in the supplied scorer.")
    add_figure(doc, "03-reward-calibration.png", "Figure 3  Supplied and selected rewards on the 200-scenario calibration split", 6.25)
    add_table(
        doc,
        ["Parameter", "Selected value", "Interpretation"],
        [
            ["Coverage count", "0.092803", "Part of mostly constant protection intercept"],
            ["Insured amount", "0.410271", "Rarely effective after normalization floor"],
            ["Deductible", "0.069603", "Part of constant protection intercept"],
            ["Coverage score", "0.427323", "Primary quality signal"],
            ["Price penalty", "1.793612", "Quality versus price trade-off"],
        ],
        widths=[1.5, 1.25, 3.8],
    )
    doc.add_heading("Falsification and held-out checks", level=2)
    add_bullets(doc, [
        "Nested five-fold selection returned 94% aggregate Top-1 and selected the same candidate in every fold.",
        "Permuted labels produced 20% nested Top-1 against a 16.3% random expectation and 49.9% pairwise accuracy.",
        "On the untouched 90-scenario silver validation split, frozen-reward Top-1 was 93.3% versus 46.7% for the default.",
        "After weights were frozen, locked golden evaluation produced 84% Fable Top-1 versus 38% for the default reward argmax.",
    ])
    add_body(doc, "The 84% scorer result is zero-inference-cost and interpretable, but it is catalogue-specific and depends on structured coverage scores. It does not replace the requested learned-model deliverable.")

    section_intro(doc, "7", "Model selection and RunPod execution", "Qwen3.5-2B was chosen as a deliberately small model for a narrow ranking policy, then trained with BF16 LoRA after staged compatibility checks.")
    doc.add_heading("Why the 2B model", level=2)
    add_bullets(doc, [
        "The prompt is text-only and the required completion is a short ranking, so broad generative capacity is less important than consistent task execution.",
        "The BF16 base checkpoint is approximately 4.55 GB and supports economical single-GPU LoRA training and serving.",
        "A 2B result creates a stronger size and cost comparison than reproducing the supplied 9B reference.",
        "BF16 LoRA avoided the documented quantization sensitivity of this Qwen3.5 family while still training only 10.9 million parameters.",
    ])
    doc.add_heading("Execution environment", level=2)
    add_table(
        doc,
        ["Component", "Recorded value"],
        [
            ["GPU", "NVIDIA RTX A6000  48 GB class"],
            ["Precision", "BF16"],
            ["Python and CUDA", "Python 3.12.3  CUDA 12.8"],
            ["PyTorch", "2.8.0 plus cu128"],
            ["Transformers", "5.5.0"],
            ["Unsloth and TRL", "2026.9.11 and 0.24.0"],
            ["Base revision", "15852e8c16360a2fea060d615a32b45270f8a8fc"],
        ],
        widths=[2.0, 4.55],
    )
    doc.add_heading("RunPod workflow", level=2)
    add_body(doc, "The repository prepared self-contained data and training bundles locally. On RunPod, the workflow verified source and dataset hashes, resolved the package combination without replacing the existing CUDA build, loaded the pinned model, ran syntax and validation-only checks, completed a one-step SFT adapter save and reload, and only then executed the full SFT run. GRPO followed the same one-step and ten-step gate pattern. Exact commands, warnings, timings, and hashes are preserved in the RunPod journal.")
    add_body(doc, "Optional fast attention and convolution kernels were initially unavailable, so early evaluation used PyTorch fallbacks. This affected speed, not prompts, weights, or reported metrics. Later kernel installation was verified separately and did not retroactively change completed results.", style="Source Note")

    section_intro(doc, "8", "Supervised fine tuning", "SFT taught both the ranking policy and the exact response contract with only 300 labelled examples.")
    add_table(
        doc,
        ["Setting", "Value"],
        [
            ["Training and validation", "300 and 90 scenarios"],
            ["Epochs and steps", "3 epochs  21 optimizer steps"],
            ["Effective batch size", "48"],
            ["Learning rate", "2 × 10⁻⁴ with warm-up and cosine decay"],
            ["LoRA", "Rank 16 alpha 16 dropout 0"],
            ["Target modules", "q k v o gate up and down projections"],
            ["Loss", "Assistant completion only"],
            ["Maximum length", "2,048 tokens"],
            ["Seed", "3407"],
        ],
        widths=[2.0, 4.55],
    )
    add_figure(doc, "04-sft-training.png", "Figure 4  SFT loss and learning-rate schedule", 6.85)
    add_body(doc, "The run completed in 378.2 seconds with 12.6 GB peak allocated GPU memory. Training loss was 0.0493. Validation loss declined at every epoch checkpoint and reached 0.0267 at step 21, which the trainer recorded as the best checkpoint. The final adapter contains BF16 LoRA tensors and preserves the frozen base model.")
    doc.add_heading("What SFT changed", level=2)
    add_body(doc, "The corrected base evaluation used a 700-token allowance because the initial 128-token diagnostic truncated most outputs. Even with the larger allowance, base strict JSON was 1/50 and mean completion length was 493.8 tokens. SFT emitted strict JSON on all 50 golden scenarios in 43.3 tokens on average. It also increased Top-1 from 22% to 84% and Top-3 recall from 66.7% to 86.7%.")
    add_figure(doc, "08-output-efficiency.png", "Figure 5  Completion length and exact-format compliance", 6.55)
    add_body(doc, "The supplied reward ratio fell from 0.844 to 0.812 while Fable agreement increased. This is expected because the supplied scorer rewards value for money, whereas the teacher and Fable rankings place more weight on coverage quality. The metrics answer different questions.")

    section_intro(doc, "9", "GRPO continuation and stopping decision", "GRPO was implemented as a controlled ablation from the preserved SFT adapter, with a frozen business reward and explicit structural penalties.")
    doc.add_heading("Objective", level=2)
    add_table(
        doc,
        ["Reward component", "Weight", "Purpose"],
        [
            ["Strict JSON", "0.10", "Exact output schema"],
            ["Valid tariffs", "0.10", "Distinct IDs from the offered panel"],
            ["Calibrated Top-1", "0.40", "Utility of the first recommendation"],
            ["Calibrated ranking", "0.40", "Discounted utility across ranks one to three"],
        ],
        widths=[2.1, 0.85, 3.6],
    )
    add_body(doc, "The 810 training prompts contained no teacher labels. The model sampled four completions per prompt with temperature 0.8, top-p 0.95, a 128-token completion limit, learning rate 5 × 10⁻⁶, KL coefficient 0.001, Dr. GRPO loss, and no reward-standard-deviation scaling.")
    doc.add_heading("Training behavior", level=2)
    add_figure(doc, "05-grpo-training.png", "Figure 6  Ten-step GRPO diagnostics", 6.3)
    add_body(doc, "All logged values were finite. Strict JSON and tariff validity stayed at 1.0, completion length remained near 41 to 42 tokens, KL declined from 0.092 to 0.047, and no completion clipping occurred. Training took 319.4 seconds and peaked at 6.9 GB allocated GPU memory. High reward and low within-group variance show that the SFT policy was already close to saturation under this objective.")
    doc.add_heading("Development gate", level=2)
    add_table(
        doc,
        ["Metric", "Frozen SFT", "GRPO 10", "Change"],
        [
            ["Top-1", "83/90  92.2%", "84/90  93.3%", "+1 scenario"],
            ["Top-3 recall", "87.0%", "86.3%", "−0.7 pp"],
            ["Strict JSON", "100%", "100%", "No change"],
            ["Reward ratio", "0.8561", "0.8597", "+0.0036"],
        ],
        widths=[1.65, 1.65, 1.65, 1.55],
    )
    doc.add_heading("Why training stopped", level=2)
    add_body(doc, "The ten-step gate showed that optimization worked, but the gain was narrow and the full ranking regressed. A longer run could still improve: ten steps are not evidence of an upper bound. However, the reward was fitted to silver labels, so stronger optimization could also amplify teacher or proxy errors. Once the authorized golden comparison showed no Top-1 gain, using that result to choose further training would violate the held-out protocol. The experiment therefore stopped with SFT as the supported model and GRPO as an inconclusive ablation.")

    section_intro(doc, "10", "Final evaluation and model selection", "SFT produced the decisive improvement; ten-step GRPO matched its first choice but slightly weakened the full ranking.")
    add_figure(doc, "06-model-progress.png", "Figure 7  Golden progression and separate reward-calibration result", 6.85)
    add_table(
        doc,
        ["System", "Top-1", "Top-3", "Strict JSON", "Reward ratio"],
        [
            ["Supplied reward argmax", "38%", "65.3%", "Deterministic", "1.000"],
            ["Calibrated scorer", "84%", "79.3%", "Deterministic", "0.812"],
            ["Qwen3.5 2B base", "22%", "66.7%", "2%", "0.844"],
            ["Qwen3.5 2B SFT", "84%", "86.7%", "100%", "0.812"],
            ["Qwen3.5 2B GRPO 10", "84%", "86.0%", "100%", "0.812"],
            ["Gemini 3 Flash", "80%", "88.0%", "Supplied baseline", "0.76"],
            ["GPT-5.4", "84%", "79.3%", "Supplied baseline", "0.80"],
            ["GPT-6 Sol", "84%", "86.7%", "Supplied baseline", "0.832"],
        ],
        widths=[2.0, 0.8, 0.85, 1.4, 1.1],
    )
    doc.add_heading("Selection decision", level=2)
    add_body(doc, "SFT is selected over GRPO because it ties golden Top-1 and supplied reward, has slightly higher Top-3 recall, and has the simpler provenance. It is selected over the calibrated scorer because the task requires a ranked shortlist: both reached 84% Top-1, but SFT recovered 130 of 150 Fable Top-3 positions versus 119 for the scorer, an 11-position gain. If only rank one were consumed, the symbolic scorer would remain the simpler production choice. GRPO changed only one golden Top-3 list: the first choice remained correct, but one teacher tariff was replaced. The result does not establish that GRPO is ineffective; it establishes that this short run did not improve the deliverable.")
    doc.add_heading("Comparison caution", level=2)
    add_body(doc, "The 2B base model is not the supplied 9B open baseline. Its 22% result should be used only as the pretraining control for our adapter. Frontier figures come from the supplied benchmark and may use different providers, prices, or output behavior. The cost-quality chart labels those provenance differences.")

    section_intro(doc, "11", "Serving benchmark and economics", "A measured vLLM deployment replaced the earlier parameter-scaled cost estimate with an observed utilization curve.")
    add_figure(doc, "09-serving-efficiency.png", "Figure 8  RTX 4090 concurrency throughput latency utilization and cost", 6.35)
    add_body(doc, "The selected adapter was served in BF16 through vLLM 0.29.0 on a RunPod RTX 4090 priced at $0.74 per hour. A sustained run sent 1,200 requests at each tested concurrency. At concurrency 64, all requests completed with valid Top-3 JSON, throughput reached 63.66 recommendations per second, and GPU utilization averaged 90.7%.")
    add_table(
        doc,
        ["Quantity", "Measured value"],
        [
            ["Mean prompt and completion", "1,023.3 and 41.3 tokens"],
            ["Latency", "p50 1.003 s  p95 1.105 s"],
            ["Throughput", "63.659 recommendations per second"],
            ["Saturated GPU cost", "€0.002832 per 1,000"],
            ["Cost at 25% utilization", "€0.011327 per 1,000"],
            ["Failures and invalid output", "0 of 1,200"],
        ],
        widths=[2.35, 4.2],
    )
    add_figure(doc, "07-cost-quality.png", "Figure 9  Quality versus cost using the measured 2B serving range", 6.35)
    add_body(doc, "These values cover GPU runtime at the configured hourly price. They exclude model download, cold start, storage, network, taxes, idle redundancy, orchestration, monitoring, and engineering labor. The saturated figure is a hardware-specific lower bound; the 25% utilization case is the planning point used in the report.")
    doc.add_heading("Recorded experiment spend", level=2)
    add_table(
        doc,
        ["Expense", "Recorded cost", "Scope"],
        [
            ["GPT-6 Sol teacher selection", "About $0.16", "Initial 50-scenario selection run"],
            ["Teacher batch annotation", "$3.635911", "1,490 validated silver labels"],
            ["RunPod", "$4.00", "Training evaluation and serving benchmarks"],
            ["Total", "About $7.80", "End-to-end recorded cash spend"],
        ],
        widths=[2.15, 1.25, 3.15],
    )
    add_body(doc, "The RunPod amount is the owner-reported total bill. The repository artifacts independently preserve workload runtimes and the serving benchmark's $0.74 per hour RTX 4090 rate, but do not allocate the $4.00 bill across individual GPU stages.", style="Source Note")

    section_intro(doc, "12", "Reproducibility and release", "The release contains portable adapters, non-golden prepared data, exact configurations, aggregate evaluations, and integrity metadata without duplicating the public base model.", break_before=False)
    add_bullets(doc, [
        "release/model/sft-final contains the selected LoRA adapter; release/model/grpo-10step preserves the ablation.",
        "The base model is pinned to Hugging Face revision 15852e8c16360a2fea060d615a32b45270f8a8fc and verified by weight checksum.",
        "release/data contains the 300/90 SFT and 810/90 GRPO splits plus development-evaluation records; the golden dataset is excluded.",
        "MANIFEST.json and SHA256SUMS record released files; scripts/create_release_bundle.py verifies the allowlist and rejects sensitive content.",
        "scripts/reproduce_model.py verifies imported artifacts, reproduces SFT, optionally runs ten-step GRPO, and uses development evaluation by default.",
        "HOW_TO_RUN.md is the reviewer entry point for Git LFS, environment creation, evaluation, retraining, and expected metrics.",
    ])
    doc.add_heading("Cross-platform verification", level=2)
    add_body(doc, "A portable Transformers backend evaluated the released SFT adapter on an Apple M4 MacBook Air using MPS. All 50 outputs matched the canonical CUDA evaluation exactly on raw text, ranked IDs, correctness, validity, rewards, and token counts. Only timing differed. This verifies that the adapter and prompt contract are not tied to the original RunPod environment.")
    add_table(
        doc,
        ["Integrity item", "Value"],
        [
            ["SFT adapter directory SHA-256", "f3bbd9330f353c038ce3d9e2338c76b6f5cd4b036ee81cd23bc7ed510c66bf48"],
            ["GRPO adapter directory SHA-256", "474543cab81b246c90b2b7c16bf79f714e7e4cdbedbba5c5826269106be25954"],
            ["Golden dataset SHA-256", "864123d2f1c2741f94cfef94ebdd09ecc6eb6d1b3fe898212441ad99b54eaa33"],
        ],
        widths=[2.1, 4.45],
    )

    section_intro(doc, "13", "Limitations and future experiments", "The reported result is strong inside the supplied digital twin, but the small benchmark, proxy supervision, and narrow product catalogue limit the claim.")
    doc.add_heading("Limitations", level=2)
    add_table(
        doc,
        ["Risk", "Why it matters", "Mitigation"],
        [
            ["Silver labels", "GPT-6 Sol differs from Fable on 8 of 50 first choices", "Use broker adjudication or a teacher ensemble"],
            ["Small golden set", "One scenario moves Top-1 by two percentage points", "Create a larger untouched expert test set"],
            ["Catalogue repetition", "Only nine products may permit catalogue-specific shortcuts", "Test new insurers and tariff revisions"],
            ["Synthetic risks", "Generated marginals may miss real correlations", "Evaluate on production-like historical risks"],
            ["Quote feasibility", "FHG3 and other failures are filtered from training", "Measure and report selection bias"],
            ["Reward proxy", "Longer GRPO may amplify proxy errors", "Use expert-quality reward development data"],
            ["Cost scope", "GPU-only estimates omit operational overhead", "Benchmark the production deployment stack"],
        ],
        widths=[1.25, 2.75, 2.55],
    )
    doc.add_heading("Predeclared next experiment", level=2)
    add_body(doc, "Collect a new broker-labelled development set and a separate untouched test set. Mine cases where Fable, GPT-6 Sol, the calibrated scorer, and SFT disagree. Train either a preference model or revised reward on the development labels, freeze it, then run a checkpointed GRPO schedule such as 10, 20, 30, 40, and 50 steps. Select the earliest checkpoint that improves Top-1 without reducing Top-3 or format validity on the new development set. Evaluate once on the new test set.")
    doc.add_heading("Additional product work", level=2)
    add_bullets(doc, [
        "Add richer policy-wording and exclusion features so coverage quality is not compressed into one score.",
        "Test out-of-catalogue insurers and shifted coverage-score definitions.",
        "Calibrate confidence and define a human-review route for low-margin rankings.",
        "Measure arrival rates, batching, autoscaling, redundancy, and cold-start behavior in the target deployment.",
        "Monitor preference drift as premiums and tariff definitions change.",
    ])

    section_intro(doc, "Appendix A", "Decision trail", "The final design emerged through explicit gates rather than a single uninterrupted training run.")
    add_table(
        doc,
        ["Decision", "Choice", "Evidence"],
        [
            ["Training data", "Generate non-golden risks and use live quotes", "Held-out rule and real workflow requirement"],
            ["Postcodes", "Use known-valid supplied postcodes", "Pilot exposed fabricated-location risk"],
            ["Cache semantics", "Reuse every cached status", "Protect live-call budget and resumability"],
            ["Teacher", "GPT-6 Sol silver labels", "84% golden Top-1 and scalable structured output"],
            ["Reward", "Freeze candidate 15957", "Nested CV stability and permutation control"],
            ["Base model", "Qwen3.5-2B BF16 LoRA", "Narrow task and inexpensive deployment target"],
            ["SFT checkpoint", "Final step 21 adapter", "Lowest validation loss"],
            ["GRPO gate", "One step then ten steps", "Compatibility and early regression detection"],
            ["Final model", "SFT adapter", "Same Top-1 as GRPO with higher Top-3 recall"],
            ["Cost claim", "Measured utilization curve", "RTX 4090 vLLM benchmark"],
            ["Submission", "Four-page brief plus technical report", "Preserve page limit and full evidence"],
        ],
        widths=[1.45, 2.25, 2.85],
    )
    add_body(doc, "Full options, rationale, and reversal conditions are recorded in DECISIONS.md.", style="Source Note")

    section_intro(doc, "Appendix B", "Evidence map and open questions", "Exact claims should be traced to frozen records; owner input is reserved for framing choices that the artifacts cannot answer.")
    add_table(
        doc,
        ["Topic", "Primary repository evidence"],
        [
            ["Narrative chronology", "gpt.md and documentation/runpod_training_journal.md"],
            ["Experiment ledger", "documentation/experiment-runs.json"],
            ["Dataset", "artifacts/quoted/final_distribution_summary.json and split_distribution_report.md"],
            ["Teacher", "artifacts/teacher/annotation_summary.json"],
            ["Reward", "artifacts/reward_calibration calibration and validation reports"],
            ["SFT and GRPO", "artifacts/training run records and release evaluation summaries"],
            ["Serving", "artifacts/serving_benchmark summaries telemetry and README"],
            ["Reproduction", "HOW_TO_RUN.md release/README.md and documentation/macos-verification.md"],
            ["Decision rationale", "DECISIONS.md and documentation/training-plan.md"],
        ],
        widths=[1.75, 4.8],
    )
    doc.add_heading("Owner wording to confirm", level=2)
    add_bullets(doc, [
        "Whether the report should call Fable an expert reference, a frontier-model reference, or both.",
        "Whether to identify the author and include repository or Hugging Face links on the title page.",
        "Whether the supplementary report will be uploaded with the brief or linked from the repository only.",
        "Whether the reader should see all internal golden-evaluation chronology or only the frozen final comparisons and held-out safeguards.",
    ])
    doc.add_heading("Principal sources", level=2)
    for source in [
        "README.md and BENCHMARK.md",
        "gpt.md",
        "DECISIONS.md",
        "documentation/runpod_training_journal.md",
        "documentation/experiment-runs.json",
        "documentation/training-plan.md",
        "documentation/file-mapping.md",
        "documentation/macos-verification.md",
        "artifacts/quoted/final_distribution_report.md",
        "artifacts/reward_calibration/calibration_report.md",
        "artifacts/reward_calibration/validation/validation_report.md",
        "release/evaluation/*.json",
        "artifacts/serving_benchmark/*/summary.json",
        "HOW_TO_RUN.md and release/README.md",
    ]:
        add_body(doc, source, style="Small Body")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--seed", type=int, default=42, help="Reserved for deterministic report generation")
    parser.add_argument("--output-dir", type=Path, default=REPORTS)
    args = parser.parse_args()
    _ = args.seed
    build_brief(args.output_dir / "wg_recommendation_submission_brief.docx")
    build_technical(args.output_dir / "wg_recommendation_technical_report.docx")


if __name__ == "__main__":
    main()
